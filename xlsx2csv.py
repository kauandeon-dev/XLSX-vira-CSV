#!/usr/bin/env python3
"""
Conversor em lote de .xlsx/.xlsm para .csv

"""

import argparse
import csv
import re
import sys
from datetime import datetime, date, time
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Falta a dependência: pip install openpyxl")


def slug(nome: str) -> str:
    return re.sub(r"[^\w\-.]+", "_", nome).strip("_") or "sheet"


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S") if (v.hour or v.minute or v.second) else v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, bool):
        return "1" if v else "0"
    return str(v)


def dump_sheet(ws, destino: Path, delim: str, encoding: str, quoting) -> int:
    linhas = 0
    with destino.open("w", newline="", encoding=encoding, errors="replace") as fh:
        w = csv.writer(fh, delimiter=delim, quoting=quoting, lineterminator="\n")
        for row in ws.iter_rows(values_only=True):
            if row is None or all(c is None for c in row):
                continue
            w.writerow([fmt(c) for c in row])
            linhas += 1
    return linhas


def main():
    p = argparse.ArgumentParser(description="Converte .xlsx/.xlsm em .csv em lote")
    p.add_argument("entrada", nargs="?", default=".", help="Pasta (ou arquivo) de origem")
    p.add_argument("-o", "--saida", help="Pasta de destino (padrão: mesma da origem)")
    p.add_argument("-r", "--recursivo", action="store_true", help="Varre subpastas")
    p.add_argument("-d", "--delim", default=",", help="Delimitador (padrão: ,)")
    p.add_argument("-e", "--encoding", default="utf-8-sig", help="Encoding (padrão: utf-8-sig)")
    p.add_argument("--all-sheets", action="store_true", help="Exporta todas as abas, não só a primeira")
    p.add_argument("--quote-all", action="store_true", help="Coloca aspas em todos os campos")
    p.add_argument("--overwrite", action="store_true", help="Sobrescreve csv já existente")
    args = p.parse_args()

    quoting = csv.QUOTE_ALL if args.quote_all else csv.QUOTE_MINIMAL
    entrada = Path(args.entrada).expanduser().resolve()

    if entrada.is_file():
        arquivos = [entrada]
        base = entrada.parent
    else:
        if not entrada.is_dir():
            sys.exit(f"Caminho inválido: {entrada}")
        padrao = "**/*" if args.recursivo else "*"
        arquivos = sorted(
            f for f in entrada.glob(padrao)
            if f.is_file()
            and f.suffix.lower() in (".xlsx", ".xlsm")
            and not f.name.startswith("~$")
        )
        base = entrada

    if not arquivos:
        sys.exit("Nenhum .xlsx encontrado.")

    saida_base = Path(args.saida).expanduser().resolve() if args.saida else None
    if saida_base:
        saida_base.mkdir(parents=True, exist_ok=True)

    ok = falhas = pulados = 0

    for arq in arquivos:
        try:
            wb = load_workbook(arq, read_only=True, data_only=True)
        except Exception as e:
            print(f"[ERRO] {arq.name}: {e}")
            falhas += 1
            continue

        if saida_base:
            destino_dir = saida_base / arq.parent.relative_to(base) if args.recursivo else saida_base
        else:
            destino_dir = arq.parent
        destino_dir.mkdir(parents=True, exist_ok=True)

        abas = wb.worksheets if args.all_sheets else wb.worksheets[:1]

        for ws in abas:
            nome = f"{arq.stem}__{slug(ws.title)}.csv" if args.all_sheets and len(wb.worksheets) > 1 else f"{arq.stem}.csv"
            destino = destino_dir / nome

            if destino.exists() and not args.overwrite:
                print(f"[SKIP] {destino.name} já existe (use --overwrite)")
                pulados += 1
                continue

            try:
                n = dump_sheet(ws, destino, args.delim, args.encoding, quoting)
                print(f"[OK]   {arq.name} :: {ws.title} -> {destino.name} ({n} linhas)")
                ok += 1
            except Exception as e:
                print(f"[ERRO] {arq.name} :: {ws.title}: {e}")
                falhas += 1

        wb.close()

    print(f"\nConcluído: {ok} gerado(s), {pulados} pulado(s), {falhas} falha(s).")


if __name__ == "__main__":
    main()
